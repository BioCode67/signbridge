"""과제개요_닿다_영남대학교.docx — 사무국 양식(2-2) 1장. 한글에서 열어 .hwp로 저장한다."""
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "final_report/제출본/과제개요_닿다_영남대학교.docx"
FONT = "휴먼명조"

doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Cm(21), Cm(29.7)
s.left_margin = s.right_margin = Cm(1.8)
s.top_margin = s.bottom_margin = Cm(1.5)
st = doc.styles["Normal"]
st.font.name, st.font.size = FONT, Pt(10)
st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)


def put(cell, text, bold=False, size=10, center=False):
    cell.text = ""
    for i, ln in enumerate(text.split("\n")):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.2
        if center:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(ln)
        r.bold, r.font.size, r.font.name = bold, Pt(size), FONT
        r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)


def shade(cell):
    s_ = OxmlElement("w:shd")
    s_.set(qn("w:val"), "clear")
    s_.set(qn("w:fill"), "E7EEF8")
    cell._element.get_or_add_tcPr().append(s_)


t = doc.add_table(rows=8, cols=4)
t.style = "Table Grid"
t.alignment = WD_TABLE_ALIGNMENT.CENTER
t.autofit = False
W = [Cm(2.6), Cm(7.2), Cm(2.0), Cm(5.6)]
for j, w in enumerate(W):
    t.columns[j].width = w
for row in t.rows:
    for j, w in enumerate(W):
        row.cells[j].width = w

put(t.cell(0, 0), "팀명", True, center=True); shade(t.cell(0, 0))
put(t.cell(0, 1), "닿다")
put(t.cell(0, 2), "분야", True, center=True); shade(t.cell(0, 2))
put(t.cell(0, 3), "응용서비스(N/W 활용 서비스)\nAI · 접근성")

ROWS = [
    ("과제명", "SignBridge — KOREN 기반 재난정보 한국수어(KSL) 양방향 AI 통역 서비스"),
    ("과제소개",
     "재난문자는 한국어로만 오고, 농인에게 한국수어는 어순·문법이 다른 제1언어다. SignBridge는 재난문자·창구 직원의 말을 "
     "실제 농인 수어자 동작(12,833종)으로 만든 3D 아바타 수어로 전달하고, 농인이 카메라 앞에서 한 수어를 인식(13,576종)해 "
     "한국어 문장·음성과 '가까운 대피소의 방향·거리'로 되돌려 답한다. 모든 추론이 브라우저 안에서 돌아 영상이 기기 밖으로 "
     "나가지 않으며, 통신이 끊긴 재난 현장(비행기 모드)에서도 동작한다."),
    ("KOREN\n사용목적",
     "KOREN을 Core·Edge·단말 3계층으로 나눠 쓴다. Core(수도권)에서는 AI Hub 대용량 원본을 연구망으로 받아 AI Cloud GPU로 "
     "인식·번역 모델을 학습하고 HPC로 전처리를 분산하며, KOREN VM이 24시간 원천 서버로 앱·모델을 배포한다. Edge(지역 PoP)에는 "
     "KB 단위 수어 패키지를 캐시해 재난 지역 가까이에서 송출하고, 단말은 기기 안에서 추론해 통신이 끊겨도 답한다."),
    ("KOREN 연동\n및 결과분석",
     "· AI Cloud(H200 MIG 1g.35GB): 실데이터 133GB 수신, 40만 클립 스트리밍 ETL, 모델 4종 학습·평가(에폭당 215만 세그먼트, 21~25분)\n"
     "· MIG의 '메모리 넉넉·연산 부족' 특성에 맞춰 영상 CNN 대신 2.8M 랜드마크 트랜스포머 채택 → 13,576종 분류를 학습 가능하게 함\n"
     "· 원본은 Core에 두고 밖으로는 모델(약 50MB)과 수어 패키지(문장당 약 2.4KB, 영상의 약 1/2,000)만 → 전국 동시 발령에도 회선 부담이 작음\n"
     "· KOREN VM 상시 운영 · HPC PBS 분산 변환 잡 구성 · 지역 PoP 다지점 송출 지연 실측은 후속 과제"),
]
for i, (k, v) in enumerate(ROWS, 1):
    a = t.cell(i, 0); b = t.cell(i, 1).merge(t.cell(i, 3))
    put(a, k, True, center=True); shade(a); put(b, v)

a = t.cell(5, 0); b = t.cell(5, 1).merge(t.cell(5, 3))
put(a, "네트워크\n구성도", True, center=True); shade(a)
b.text = ""
p = b.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run().add_picture(str(Path(__file__).parent / "img/koren_network.png"), width=Cm(14.4))

a = t.cell(6, 0); b = t.cell(6, 1).merge(t.cell(6, 3))
put(a, "과제\n주요 성과", True, center=True); shade(a)
put(b,
    "· 수어 낱말 인식 13,576종: 학습에 없던 수어자 13명에서 top-1 78.6% · top-5 91.9%(수어자 분리 평가)\n"
    "· 한국어→글로스 학습 모델 F1 61.3 · BLEU-4 27.1, int8 30MB로 브라우저에서 문장당 144ms\n"
    "· 번역 낱말 표현률 행동요령·길찾기 100%, 재난문자 95.8%, 창구 대화 홀드아웃 94.8% · 오역 회귀 156건 자동 검사\n"
    "· 전국 28,125곳(대피소 14,134곳) 오프라인 길 안내 · 폰·태블릿·키오스크 실조작 검증 · 복지관 현장 시연 및 당사자 평가")
t._tbl.remove(t.rows[7]._tr)
doc.save(OUT)
print("saved", OUT)
