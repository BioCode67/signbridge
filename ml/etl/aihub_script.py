"""AI Hub「재난안전정보 수어영상」**수어스크립트(xlsx)** → 한국어↔글로스 쌍.

    python -m ml.etl.aihub_script --input <수어스크립트_TL.zip> --out <데이터>/script

키포인트가 필요 없다. 372MB짜리 수어스크립트만으로 ③단계
(`ml/train_gloss2text.py --direction text2gloss`)를 학습할 수 있다 —
90GB 형태소 JSON을 기다릴 필요가 없고, 현재 앱의 규칙 기반 `signAgent.ts`를
데이터 학습 모델로 대체하는 경로다.

────────────────────────────────────────────────────────────────────────────
확인된 xlsx 구조 (실제 파일을 열어 확인. 문서에 없다)
────────────────────────────────────────────────────────────────────────────
파일 하나에 클립 **여러 개**가 세로로 쌓여 있다(한 파일 15만 행 · 클립 2,697개인 것도 있다).
각 클립 블록은 `Information` 행에서 시작한다. 열은 **시간축**이다 — 3행의 `Second :`가
1,2,3… 초를 가리킨다.

    열번호 →              0                  1              2       3       4
    ────────────────────────────────────────────────────────────────────────────
    r1   Information       File name :        NIA_SL_G1_COLDWAVE000010_1_TW07.js
    r2                     Korean sentence :  오늘 21시부로 한파가 예상되오니 …
    r3                     Second :           1       2       3
    r19  sign_gestures_both gloss_id :        오늘1
    r20                    start(s) :         1.823
    r21                    end(s) :                   2.275          ← **3열**
    r22  sign_gestures_both gloss_id :                밤1     시:9시
    r23                    start(s) :                 2.385   3.153
    r24                    end(s) :                   2.878   3.55

**가장 헷갈리는 지점**: 값은 *자기 시각*에 해당하는 열에 놓인다. 그래서 같은 항목인데도
`오늘1`(2열) · start `1.823`(2열) · end `2.275`(**3열**)로 열이 어긋난다.
열을 기준으로 짝지으면 끝 시각을 통째로 놓친다(조용히 start와 같은 값이 되어버린다).

열은 시간축 눈금일 뿐이고, **짝은 왼쪽부터 나온 순서로 맞춘다.** 즉 gloss 행의 k번째
값과 start 행의 k번째, end 행의 k번째가 한 쌍이다. 가정이 깨지면 `MISMATCH`로 보고한다.

글로스는 형태소 JSON과 마찬가지로 `both`/`strong`/`weak` 세 층렬로 나뉘므로 합쳐서
시간순 정렬한다. 비수지 층(Ci·EBf·Mmo·Mctr·Hno·Mo1·Hs·Tbt)은 글로스가 아니라 건너뛴다.

────────────────────────────────────────────────────────────────────────────
분할은 **한국어 문장 기준**이 기본이다
────────────────────────────────────────────────────────────────────────────
같은 문장을 여러 수어자가 1:1~1:3으로 연기하므로, 무작위로 나누면 **같은 문장이 학습과
평가 양쪽에 들어간다.** 그러면 번역을 배운 게 아니라 외운 것을 평가하게 된다.
`--split-by text`(기본)는 문장 해시로 갈라 이 누수를 막는다.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import re
import sys
from collections import Counter
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from ml.etl.aihub_disaster import parse_disaster_name  # noqa: E402
from ml.signbridge.archive import find_archives, stream_archive  # noqa: E402

GLOSS_TIERS = {
    "sign_gestures_both": "both",
    "sign_gestures_strong": "strong",
    "sign_gestures_weak": "weak",
}

INFO_ROW = "Information"
FILE_LABEL = "File name"
KOREAN_LABEL = "Korean sentence"
GLOSS_LABEL = "gloss_id"
START_LABEL = "start(s)"
END_LABEL = "end(s)"

# 짝 개수가 어긋난 행 수. 구조 가정이 깨지면 조용히 틀리는 대신 마지막에 보고한다.
MISMATCH: Counter = Counter()


def _cell(value) -> str:
    return "" if value is None else str(value).strip()


def _label(value) -> str:
    """`File name :` → `File name`. 배포본마다 콜론·공백이 들쭉날쭉하다."""
    return _cell(value).rstrip(":").strip()


def _to_float(value) -> float | None:
    text = _cell(value)
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def parse_sheet(rows: list[tuple], source_name: str, tiers: set[str]) -> list[dict]:
    """시트 한 장에서 클립 레코드들을 뽑는다.

    rows는 (A열, B열, C열…) 튜플의 리스트다. `Information` 행마다 새 클립이 시작된다.
    """
    clips: list[dict] = []
    current: dict | None = None

    index = 0
    while index < len(rows):
        row = rows[index]
        col_a = _label(row[0]) if len(row) > 0 else ""
        col_b = _label(row[1]) if len(row) > 1 else ""

        if col_a == INFO_ROW or col_b == FILE_LABEL:
            if current and current["glosses"]:
                clips.append(current)
            # `NIA_SL_G1_COLDWAVE000010_1_TW07.js` → 확장자를 떼어 clip id로 쓴다.
            raw_id = _cell(row[2]) if len(row) > 2 else ""
            clip_id = re.sub(r"\.(js|json|mp4)$", "", raw_id, flags=re.IGNORECASE)
            current = {
                "id": clip_id or f"{source_name}#{index}",
                "korean_text": "",
                "glosses": [],
            }
            index += 1
            continue

        if current is None:
            index += 1
            continue

        if col_b == KOREAN_LABEL:
            current["korean_text"] = _cell(row[2]) if len(row) > 2 else ""
            index += 1
            continue

        tier = GLOSS_TIERS.get(_cell(row[0]))
        if tier and col_b == GLOSS_LABEL and tier in tiers:
            start_row = rows[index + 1] if index + 1 < len(rows) else ()
            end_row = rows[index + 2] if index + 2 < len(rows) else ()
            if _label(start_row[1] if len(start_row) > 1 else "") != START_LABEL:
                start_row = ()
            if _label(end_row[1] if len(end_row) > 1 else "") != END_LABEL:
                end_row = ()

            # **열 위치로 짝지으면 안 된다.** 각 값은 자기 시각에 해당하는 열에 놓이므로
            # 같은 항목이라도 gloss·start·end가 서로 다른 열에 있다.
            #   오늘1(2열) / start 1.823(2열) / end 2.275(**3열**)
            # 즉 열은 시간축 눈금일 뿐이고, 짝은 **왼쪽부터 나온 순서**로 맞아떨어진다.
            glosses = [(c, _cell(v)) for c, v in enumerate(row[2:], 2) if _cell(v)]
            starts = [_to_float(v) for v in start_row[2:] if _cell(v)]
            ends = [_to_float(v) for v in end_row[2:] if _cell(v)]

            if len(starts) != len(glosses):
                MISMATCH["start"] += 1
            for order, (_, gloss) in enumerate(glosses):
                start = starts[order] if order < len(starts) else None
                end = ends[order] if order < len(ends) else None
                if start is None:
                    continue
                current["glosses"].append(
                    {
                        "gloss": gloss,
                        "start": start,
                        "end": end if end is not None and end >= start else start,
                        "tier": tier,
                    }
                )
            index += 3 if start_row and end_row else 1
            continue

        index += 1

    if current and current["glosses"]:
        clips.append(current)

    for clip in clips:
        # 층렬을 합쳤으므로 시간순으로 정렬해야 한 줄의 글로스열이 된다.
        clip["glosses"].sort(key=lambda g: (g["start"], g["end"]))
    return clips


def read_xlsx(raw: bytes, source_name: str, tiers: set[str]) -> list[dict]:
    import openpyxl

    workbook = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    clips: list[dict] = []
    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        rows = list(sheet.iter_rows(values_only=True))
        clips.extend(parse_sheet(rows, source_name, tiers))
    workbook.close()
    return clips


def assign_split(clip: dict, mode: str, val_ratio: float, test_ratio: float) -> str:
    """분할 키를 해시해 결정적으로 나눈다(같은 입력이면 늘 같은 분할).

    기본은 **한국어 문장** 기준이다. 같은 문장을 여러 수어자가 연기하므로 클립 단위로
    나누면 같은 문장이 학습·평가 양쪽에 들어가 성능이 부풀려진다.
    """
    if mode == "none":
        return "train"
    key = {
        "text": clip["korean_text"],
        "clip": clip["id"],
        "category": clip.get("category", ""),
    }.get(mode, clip["korean_text"])
    if not key:
        key = clip["id"]
    digest = hashlib.sha1(key.encode("utf-8")).hexdigest()
    bucket = int(digest[:8], 16) / 0xFFFFFFFF
    if bucket < test_ratio:
        return "test"
    if bucket < test_ratio + val_ratio:
        return "val"
    return "train"


def main() -> None:
    parser = argparse.ArgumentParser(
        description="수어스크립트 xlsx → 한국어↔글로스 index.jsonl (키포인트 불필요)"
    )
    parser.add_argument(
        "--input", type=Path, required=True, help="xlsx가 든 아카이브·디렉터리 (압축 그대로)"
    )
    parser.add_argument("--out", type=Path, required=True)
    parser.add_argument("--tiers", default="both,strong,weak")
    parser.add_argument(
        "--split-by",
        choices=["text", "clip", "category", "none"],
        default="text",
        help="기본 text — 같은 문장이 학습·평가에 함께 들어가는 누수를 막는다",
    )
    parser.add_argument("--val-ratio", type=float, default=0.1)
    parser.add_argument("--test-ratio", type=float, default=0.1)
    parser.add_argument("--limit", type=int, default=0, help="xlsx 몇 개만 (0=전부)")
    args = parser.parse_args()

    tiers = {t.strip() for t in args.tiers.split(",") if t.strip()}
    args.out.mkdir(parents=True, exist_ok=True)

    clips: list[dict] = []
    seen_files = 0

    def handle(name: str, raw: bytes) -> None:
        nonlocal seen_files
        if args.limit and seen_files >= args.limit:
            return
        seen_files += 1
        try:
            found = read_xlsx(raw, Path(name).stem, tiers)
        except Exception as error:
            print(f"  [warn] {name}: {type(error).__name__}: {error}", flush=True)
            return
        for clip in found:
            info = parse_disaster_name(clip["id"])
            clip["category"] = info.get("category", "")
            clip["source_file"] = name
        clips.extend(found)
        print(f"  [{seen_files}] {Path(name).name} → 클립 {len(found)}개", flush=True)

    archives = find_archives(args.input)
    if archives:
        for archive in archives:
            print(f"[script] 아카이브 스트리밍: {archive.name}")
            stream_archive(archive, (".xlsx",), handle)
    elif args.input.is_dir():
        for path in sorted(args.input.rglob("*.xlsx")):
            handle(str(path), path.read_bytes())
    else:
        handle(str(args.input), args.input.read_bytes())

    if not clips:
        raise SystemExit(f"클립을 찾지 못했습니다: {args.input}")

    # 같은 클립 id가 여러 파일에 중복될 수 있다. 글로스가 더 많은 쪽을 남긴다.
    best: dict[str, dict] = {}
    for clip in clips:
        prior = best.get(clip["id"])
        if prior is None or len(clip["glosses"]) > len(prior["glosses"]):
            best[clip["id"]] = clip
    records = sorted(best.values(), key=lambda c: c["id"])

    kept = 0
    counts = Counter()
    index_path = args.out / "index.jsonl"
    with open(index_path, "w", encoding="utf-8") as handle_out:
        for clip in records:
            if not clip["korean_text"] or not clip["glosses"]:
                continue
            split = assign_split(clip, args.split_by, args.val_ratio, args.test_ratio)
            counts[split] += 1
            kept += 1
            handle_out.write(
                json.dumps(
                    {
                        "id": clip["id"],
                        # 키포인트가 없는 텍스트 전용 레코드다. read_index가 npz 키를
                        # 요구하므로 빈 문자열을 둔다(gloss2text는 팩을 읽지 않는다).
                        "npz": "",
                        "fps": 30.0,
                        "num_frames": 0,
                        "korean_text": clip["korean_text"],
                        "glosses": clip["glosses"],
                        "signer": "",
                        "category": clip["category"],
                        "split": split,
                        "source": "aihub-disaster-script",
                    },
                    ensure_ascii=False,
                )
                + "\n"
            )

    vocab = Counter(g["gloss"] for c in records for g in c["glosses"])
    sentences = {c["korean_text"] for c in records if c["korean_text"]}
    print(f"\n[script] xlsx {seen_files}개 → 클립 {kept}개 → {index_path}")
    print(f"[script] 서로 다른 한국어 문장 {len(sentences):,}개 · 글로스 어휘 {len(vocab):,}종")
    print(f"[script] 분할({args.split_by}): " + " · ".join(f"{k} {v}" for k, v in counts.items()))
    print(f"[script] 상위 글로스: {[g for g, _ in vocab.most_common(12)]}")
    if MISMATCH:
        print(f"[script] ⚠️ gloss/start 개수가 어긋난 행 {MISMATCH['start']}개 — 구조 가정 확인 필요")

    stats = {
        "clips": kept,
        "sentences": len(sentences),
        "gloss_types": len(vocab),
        "split_by": args.split_by,
        "splits": dict(counts),
    }
    (args.out / "stats.json").write_text(
        json.dumps(stats, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(f"[script] 통계 → {args.out / 'stats.json'}")


if __name__ == "__main__":
    main()
