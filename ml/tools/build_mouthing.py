"""마우징 표를 만든다 — 글로스 → 그때 입으로 내는 한국어 낱말.

    python3 ml/tools/build_mouthing.py            # public/data/mouthing.json
    python3 ml/tools/build_mouthing.py --report   # 만들지 않고 통계만

**마우징이 무엇인가.** 수어를 하면서 입으로 한국어 낱말을 소리 없이 발음하는 것이다.
같은 손동작이 여러 뜻을 가질 때 입모양이 뜻을 가른다 — 수어의 일부지 장식이 아니다.

**원본.** AI Hub 수어스크립트 엑셀(`~/sbdata/nms-script`)의 `Mmo` 채널.
131개 파일에 29,352건 · 낱말 3,435종이 시간 구간까지 붙어 있다(2026-08-20 실측).
같은 파일의 8채널 중 **descriptor(무엇을)가 채워진 것은 Mmo뿐**이다 —
눈썹(EBf) 7건, 고개(Hno) 10건뿐이라 "언제 움직이는지"만 알고 "어떻게"는 모른다.
그래서 눈썹·고개는 여기서 만들지 않는다.

**어떻게 잇나.** Mmo 구간과 글로스 구간의 **겹치는 시간**이 가장 긴 짝을 고른다.
한 글로스에 여러 마우징이 붙으면 가장 자주 나온 것을 쓴다.
"""
from __future__ import annotations
import argparse, glob, json, os, re
from collections import Counter, defaultdict
from pathlib import Path

NMS = Path.home() / "sbdata" / "nms-script"
OUT = Path("public/data/mouthing.json")
# 손동작 층렬 — 셋을 합쳐 시간순으로 본다(CLAUDE.md의 함정 표에 있는 그대로).
TIERS = ("sign_gestures_both", "sign_gestures_strong", "sign_gestures_weak")


def rows_of(path: str):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    out = list(ws.iter_rows(values_only=True))
    wb.close()
    return out


def parse(path: str):
    """한 파일에서 (글로스 구간들, 마우징 구간들)을 뽑는다.

    **열 번호로 짝지으면 안 된다.** 엑셀은 값을 "몇 초쯤"에 해당하는 칸에
    적어 두는데, 라벨·start·end 세 줄의 칸이 서로 어긋나 있다:

        18  [0]sign_gestures_both  [1]gloss_id :  [2]오늘1
        19                         [1]start(s) :  [2]1.823
        20                         [1]end(s) :    [3]2.275      ← 칸이 하나 밀렸다

    **비어 있지 않은 값을 순서대로** 세어 짝지어야 한다. 칸으로 맞추면
    같은 낱말이 여러 번 세어지고(실측: 글로스 구간이 163만 개로 부풀었다)
    마우징이 아무 낱말에나 달라붙는다(전부 `중`으로 나왔다).
    """
    rows = rows_of(path)
    # **문장마다 따로 담는다.** 파일 하나에 문장이 수백 개 들어 있고 시간은
    # 문장마다 0부터 다시 센다. 파일 전체를 한 덩어리로 보면 다른 문장의
    # 마우징이 겹쳐 들어와 아무 낱말에나 지명이 달라붙는다
    # (실측: `조심1→얼음`·`장소1→임진`·`춥다1→봉화군`).
    blocks: list[tuple[list, list]] = []
    glosses: list[tuple[str, float, float]] = []
    mouths: list[tuple[str, float, float]] = []
    cur = None
    labels: list[str] = []
    starts: list[float] = []

    def cells(r):
        return [c for c in r[2:] if c not in (None, "")]

    def nums(r):
        out = []
        for c in cells(r):
            try:
                out.append(float(c))
            except (TypeError, ValueError):
                pass
        return out

    for r in rows:
        c0 = r[0].strip() if isinstance(r[0], str) else ""
        c1 = r[1].strip() if isinstance(r[1], str) else ""
        if c1.startswith("Korean sentence"):
            blocks.append((glosses, mouths))
            glosses, mouths = [], []
            cur = None
            labels, starts = [], []
            continue
        if c0 and c0 != "Information":
            cur = c0
        if not cur:
            continue
        if c1.startswith("gloss_id") or c1.startswith("descriptor"):
            # 라벨에 개행·공백이 섞여 있다 — 안 떼면 같은 낱말이 둘로 갈린다.
            labels = [str(c).strip() for c in cells(r)]
            starts = []
        elif c1.startswith("start"):
            starts = nums(r)
        elif c1.startswith("end"):
            ends = nums(r)
            into = glosses if cur in TIERS else (mouths if cur == "Mmo" else None)
            if into is not None:
                n = min(len(labels), len(starts), len(ends))
                for i in range(n):
                    if labels[i]:
                        into.append((labels[i], starts[i], ends[i]))
            labels, starts = [], []
    blocks.append((glosses, mouths))
    return [b for b in blocks if b[0]]


def _pair(gl, mo, pair) -> None:
    """한 문장 안에서 글로스와 마우징을 **겹치는 시간이 가장 긴 짝**으로 잇는다."""
    for g, gs, ge in gl:
        best, bestOv = None, 0.0
        for m, ms, me in mo:
            ov = min(ge, me) - max(gs, ms)
            if ov > bestOv:
                bestOv, best = ov, m
        # 겹침이 글로스 길이의 30%는 돼야 같은 자리로 본다.
        if best and bestOv >= 0.30 * max(1e-6, ge - gs):
            pair[g][best] += 1


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--report", action="store_true", help="만들지 않고 통계만 낸다")
    ap.add_argument("--min-count", type=int, default=2,
                    help="이만큼은 같이 나와야 표에 넣는다(한 번뿐인 짝은 잡음일 수 있다)")
    a = ap.parse_args()

    files = sorted(glob.glob(str(NMS / "**" / "*.xlsx"), recursive=True))
    if not files:
        print(f"[마우징] 원본이 없습니다 — {NMS}")
        return 1

    pair: dict[str, Counter] = defaultdict(Counter)
    ng = nm = 0
    for i, f in enumerate(files):
        try:
            blocks = parse(f)
        except Exception as err:  # 파일 하나가 깨져도 나머지를 계속 본다
            print(f"  · 건너뜀 {os.path.basename(f)} — {err}")
            continue
        for gl, mo in blocks:
            ng += len(gl); nm += len(mo)
            _pair(gl, mo, pair)
        if (i + 1) % 25 == 0:
            print(f"  · {i + 1}/{len(files)} 파일")

    # 행정구역 이름표 — 재난문자에는 지역명이 늘 붙어 있어 아무 낱말에나
    # 달라붙는다. 글로스 자체가 지명인 경우(`임진→임진`)는 남긴다.
    places = set()
    pf = Path("ml/data/place_names.txt")
    if pf.exists():
        places = {ln.strip() for ln in pf.read_text(encoding="utf-8").splitlines() if ln.strip()}

    def hangul(w: str) -> bool:
        return any(0xAC00 <= ord(ch) <= 0xD7A3 for ch in w)

    table = {}
    dropped_num = dropped_place = dropped_time = dropped_weak = 0
    for g, c in pair.items():
        w, n = c.most_common(1)[0]
        if n < a.min_count:
            continue
        # 한글이 아니면 입모양을 만들 수 없다(숫자·기호). 넣어 두면 "마우징이
        # 있다"고 표시되는데 입은 안 움직인다 — 그게 가장 헷갈리는 상태다.
        if not hangul(w):
            dropped_num += 1
            continue
        lemma = re.sub(r"[0-9#]+$", "", g)
        if w in places and lemma not in places:
            dropped_place += 1
            continue
        # **시각·날짜는 표에 넣지 않는다.** 주석자가 `시:3시50분 → 3시50분`처럼
        # 숫자를 그대로 적어 두어 입모양을 만들 수 없고(한글은 `시`·`분`뿐),
        # 잡음도 섞여 있다(`시:21시00분 → 아홉시`·`시:13시 → 한시`).
        # 앱은 국어 수 읽기 규칙(`readableGloss`)으로 옮기는 편이 정확하다.
        if g.startswith(("시:", "날짜:")):
            dropped_time += 1
            continue
        # **이름과 다른 마우징은 더 엄격히 본다.**
        #
        # 마우징이 글로스 이름과 같으면(`강1→강`) 틀릴 여지가 없다. 다른 것은
        # 진짜 마우징 변이(`운전1→차`·`눈내리다1→눈`)일 수도, 통계 잡음
        # (`지시1#→앞`·`키우다→곳`·`인터넷→중`)일 수도 있다. 표본으로 보니
        # 반반이었다. **잡음이면 입이 엉뚱한 낱말을 말한다** — 이름을 그대로
        # 쓰는 되돌림(실측 73.8% 정확)보다 나쁘다. 그래서 다른 값은
        # 5회 이상 · 우세율 60% 이상일 때만 받는다.
        total = sum(c.values())
        if w != lemma and (n < 5 or n / total < 0.6):
            dropped_weak += 1
            continue
        table[g] = w
    # **표제어 열쇠도 함께 싣는다.** 표는 `오늘1`처럼 이형태 번호가 붙은 이름으로
    # 되어 있는데, 번역이 내는 글로스는 `오늘`처럼 번호가 없을 때가 있다. 그대로
    # 두면 입모양이 안 붙는다 — 실측에서 수록률이 27%에 머물렀고 빠진 것 대부분이
    # 이 경우였다(`안전1`·`대피0`·`오늘`·`지역`).
    lemma_votes: dict[str, Counter] = defaultdict(Counter)
    for g, w in table.items():
        lemma_votes[re.sub(r"[0-9#]+$", "", g)][w] += sum(pair[g].values())
    added = 0
    for lem, c in lemma_votes.items():
        if lem and lem not in table:
            table[lem] = c.most_common(1)[0][0]
            added += 1
    print(f"[마우징] 표제어 열쇠 {added:,}개 추가")

    print(f"[마우징] 파일 {len(files)} · 글로스 구간 {ng:,} · 마우징 구간 {nm:,}")
    print(f"[마우징] 짝지어진 글로스 {len(pair):,}종 → {a.min_count}회 이상 {len(table):,}종"
          f" (한글 아님 {dropped_num} · 지명 {dropped_place} · 시각·날짜 {dropped_time}"
          f" · 근거 약함 {dropped_weak} 제외)")
    if pair:
        ex = sorted(pair.items(), key=lambda kv: -sum(kv[1].values()))[:10]
        print("  상위:", " · ".join(f"{g}→{c.most_common(1)[0][0]}({sum(c.values())})" for g, c in ex))
    if a.report:
        return 0
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(table, ensure_ascii=False), encoding="utf-8")
    print(f"[마우징] 저장 {OUT} ({OUT.stat().st_size / 1024:.0f}KB)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
